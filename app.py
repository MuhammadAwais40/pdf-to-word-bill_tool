import io
import os
import re
import sys
import json
import time
import uuid
import math
import tempfile
import statistics
from collections import Counter
from pathlib import Path
from flask import Flask, render_template, request, jsonify, send_file, send_from_directory
import pdfplumber
from pdfminer.pdftypes import resolve1
from PIL import Image as PILImage
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
def resource_path(rel):
    base = Path(getattr(sys, "_MEIPASS", Path(__file__).resolve().parent))
    return base / rel
BASE_DIR = Path(__file__).resolve().parent
if getattr(sys, "frozen", False):
    APP_DATA = Path(os.environ.get("LOCALAPPDATA", tempfile.gettempdir())) / "PDFConverter"
else:
    APP_DATA = BASE_DIR
STORAGE = APP_DATA / "storage"
STORAGE.mkdir(parents=True, exist_ok=True)
app = Flask(
    __name__,
    template_folder=str(resource_path("templates")),
    static_folder=str(resource_path("static")),)
app.config["MAX_CONTENT_LENGTH"] = 50 * 1024 * 1024
CTRL_RE = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")
BULLET_RE = re.compile(r"^(\u2022|\u25CF|\u25AA|-|\*|\d+[\.\)]|\(?[a-zA-Z]\)|[a-zA-Z]\.)\s+")
FID_RE = re.compile(r"^[0-9a-f]{32}$")
MIME_DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
MIME_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
# ---- Image settings ----
MIN_IMG_PT = 6          
MAX_IMG_PX = 1600       
MAX_IMAGES = 300        
MAX_DOCX_WIDTH_IN = 6.5  
# ------------------------------------------------------------------
# Helpers
# ------------------------------------------------------------------
def clean_text(s):
    s = CTRL_RE.sub("", s or "")
    return re.sub(r"\s+", " ", s).strip()
def clean_cell(s):
    s = CTRL_RE.sub("", s or "")
    lines = [re.sub(r"[ \t]+", " ", ln).strip() for ln in s.splitlines()]
    return "\n".join(ln for ln in lines if ln)
def normalize_rows(rows):
    rows = [[clean_cell(c) for c in r] for r in (rows or []) if r]
    rows = [r for r in rows if any(r)]
    if not rows:
        return []
    ncols = max(len(r) for r in rows)
    rows = [r + [""] * (ncols - len(r)) for r in rows]
    keep = [i for i in range(ncols) if any(r[i] for r in rows)]
    return [[r[i] for i in keep] for r in rows]
def looks_like_table(rows):
    if len(rows) < 3:
        return False
    ncols = len(rows[0])
    if ncols < 2 or ncols > 12:
        return False
    total = len(rows) * ncols
    filled = sum(1 for r in rows for c in r if c)
    multi = sum(1 for r in rows if sum(1 for c in r if c) >= 2)
    return filled / total >= 0.5 and multi / len(rows) >= 0.6
def find_tables(page):
    found = []
    for t in page.find_tables():
        rows = normalize_rows(t.extract())
        if rows and len(rows[0]) >= 2:
            found.append((t.bbox, rows))
    if found:
        return found
    try:
        settings = {
            "vertical_strategy": "text",
            "horizontal_strategy": "text",
            "snap_tolerance": 3,
            "join_tolerance": 3,
            "intersection_tolerance": 5,
            "min_words_vertical": 3,
            "min_words_horizontal": 2,
        }
        for t in page.find_tables(settings):
            rows = normalize_rows(t.extract())
            if looks_like_table(rows):
                found.append((t.bbox, rows))
    except Exception:
        pass
    return found
def make_filter(bboxes):
    def outside(obj):
        if obj.get("object_type") != "char":
            return True
        cx = (obj["x0"] + obj["x1"]) / 2
        cy = (obj["top"] + obj["bottom"]) / 2
        return not any(b[0] <= cx <= b[2] and b[1] <= cy <= b[3] for b in bboxes)
    return outside
def line_info(ln):
    text = clean_text(ln.get("text", ""))
    if not text:
        return None
    chars = [c for c in ln.get("chars", []) if c.get("text", "").strip()]
    sizes = [c.get("size", 0) for c in chars] or [0]
    size = round(statistics.median(sizes) * 2) / 2
    bold_n = sum(1 for c in chars if re.search(r"bold|black|heavy", c.get("fontname", ""), re.I))
    bold = bool(chars) and bold_n / len(chars) > 0.8
    return {
        "text": text, "size": size, "bold": bold,
        "top": float(ln["top"]), "bottom": float(ln["bottom"]),
    }
def build_paragraphs(lines, body):
    paras, cur = [], None
    for ln in lines:
        ratio = ln["size"] / body if body else 1
        level = 0
        if ratio >= 1.2:
            level = 1 if ratio >= 1.7 else (2 if ratio >= 1.4 else 3)
        split = True
        if cur:
            gap = ln["top"] - cur["bottom"]
            limit = max(ln["size"], cur["size"], 1) * 0.6
            split = bool(
                gap > limit or level or cur["heading"]
                or BULLET_RE.match(ln["text"])
                or abs(ln["size"] - cur["size"]) > 1
                or ln["bold"] != cur["bold"]
            )
        if cur and not split:
            if cur["text"].endswith("-") and ln["text"][:1].islower():
                cur["text"] = cur["text"][:-1] + ln["text"]
            else:
                cur["text"] += " " + ln["text"]
            cur["bottom"] = ln["bottom"]
        else:
            if cur:
                paras.append(cur)
            cur = {"text": ln["text"], "size": ln["size"], "bold": ln["bold"],
                   "heading": level, "bottom": ln["bottom"]}
    if cur:
        paras.append(cur)
    return [{"type": "para", "text": p["text"], "size": p["size"],
             "bold": p["bold"], "heading": p["heading"]} for p in paras]
# ------------------------------------------------------------------
# Image extraction 
# ------------------------------------------------------------------
def _names(value):
    value = resolve1(value)
    if value is None:
        return []
    if not isinstance(value, list):
        value = [value]
    out = []
    for x in value:
        x = resolve1(x)
        out.append(getattr(x, "name", None) or (x if isinstance(x, str) else ""))
    return out
def _channels(cs):
    names = _names(cs)
    if not names:
        return 0
    first = names[0]
    if first in ("DeviceRGB", "RGB", "CalRGB"):
        return 3
    if first in ("DeviceGray", "G", "CalGray"):
        return 1
    if first == "ICCBased":
        try:
            raw = resolve1(cs)
            n = resolve1(resolve1(raw[1]).attrs.get("N"))
            return n if n in (1, 3) else 0
        except Exception:
            return 0
    return 0
def _decode_embedded(img):
    stream = img.get("stream")
    if stream is None:
        return None
    filters = _names(stream.attrs.get("Filter"))
    if filters and filters[-1] in ("DCTDecode", "DCT", "JPXDecode"):
        return PILImage.open(io.BytesIO(stream.get_data()))
    unsupported = {"CCITTFaxDecode", "CCF", "JBIG2Decode"}
    if any(f in unsupported for f in filters):
        return None
    w, h = img.get("srcsize") or (0, 0)
    bits = img.get("bits") or 8
    ch = _channels(img.get("colorspace"))
    if not (w and h and ch and bits == 8):
        return None
    data = stream.get_data()
    need = w * h * ch
    if len(data) < need:
        return None
    return PILImage.frombytes("RGB" if ch == 3 else "L", (w, h), data[:need])
def _render_region(page, img):
    x0 = max(page.bbox[0], img["x0"])
    top = max(page.bbox[1], img["top"])
    x1 = min(page.bbox[2], img["x1"])
    bottom = min(page.bbox[3], img["bottom"])
    if x1 - x0 < 1 or bottom - top < 1:
        return None
    return page.crop((x0, top, x1, bottom)).to_image(resolution=150).original
def _save_image(pil, fid, n):
    if pil.mode not in ("RGB", "RGBA", "L"):
        pil = pil.convert("RGBA" if "A" in pil.mode or "transparency" in pil.info else "RGB")
    if max(pil.size) > MAX_IMG_PX:
        pil.thumbnail((MAX_IMG_PX, MAX_IMG_PX))
    name = f"{fid}_{n}.png"
    pil.save(STORAGE / name, format="PNG", optimize=True)
    return name
def extract_page_images(page, fid, counter, has_text):
    out = []
    page_area = float(page.width) * float(page.height) or 1.0
    for img in page.images:
        if counter[0] >= MAX_IMAGES:
            break
        try:
            w = float(img["x1"]) - float(img["x0"])
            h = float(img["bottom"]) - float(img["top"])
            if w < MIN_IMG_PT or h < MIN_IMG_PT:
                continue  
            if has_text and (w * h) / page_area > 0.9:
                continue  
            pil = None
            try:
                pil = _decode_embedded(img)
            except Exception:
                pil = None
            if pil is None:
                pil = _render_region(page, img)
            if pil is None:
                continue
            counter[0] += 1
            fname = _save_image(pil, fid, counter[0])
            cx = (float(img["x0"]) + float(img["x1"])) / 2
            pw = float(page.width)
            align = "left" if cx < pw * 0.38 else ("right" if cx > pw * 0.62 else "center")
            out.append((float(img["top"]), "image", {
                "type": "image", "file": fname, "width": w, "height": h, "align": align,
            }))
        except Exception:
            continue
    return out
def remove_images(fid):
    if not FID_RE.match(fid or ""):
        return
    for p in STORAGE.glob(f"{fid}_*.png"):
        try:
            p.unlink()
        except OSError:
            pass
# ------------------------------------------------------------------
# PDF -> structured data
# ------------------------------------------------------------------
def extract_pdf(path, original_name, fid=None):
    raw_pages, sizes = [], Counter()
    img_counter = [0]
    with pdfplumber.open(path) as pdf:
        for idx, page in enumerate(pdf.pages, 1):
            items = []
            try:
                tables = find_tables(page)
            except Exception:
                tables = []
            bboxes = [t[0] for t in tables]
            try:
                src = page.filter(make_filter(bboxes)) if bboxes else page
                raw_lines = src.extract_text_lines(strip=True, return_chars=True)
            except Exception:
                raw_lines = []
            for ln in raw_lines:
                d = line_info(ln)
                if d:
                    items.append((d["top"], "line", d))
                    sizes[d["size"]] += len(d["text"])
            for bbox, rows in tables:
                items.append((bbox[1], "table", rows))
            if fid:
                try:
                    has_text = any(k == "line" for _, k, _ in items)
                    items += extract_page_images(page, fid, img_counter, has_text)
                except Exception:
                    pass
            items.sort(key=lambda x: x[0])
            raw_pages.append({"number": idx, "items": items})
            try:
                page.flush_cache()
            except Exception:
                pass
    body = sizes.most_common(1)[0][0] if sizes else 11
    pages = []
    for pg in raw_pages:
        blocks, buf = [], []
        for _, kind, payload in pg["items"]:
            if kind == "line":
                buf.append(payload)
            elif kind == "table":
                blocks += build_paragraphs(buf, body)
                buf = []
                blocks.append({"type": "table", "rows": payload})
            else:  # image
                blocks += build_paragraphs(buf, body)
                buf = []
                blocks.append(payload)
        blocks += build_paragraphs(buf, body)
        pages.append({"number": pg["number"], "blocks": blocks})
    paras = sum(1 for p in pages for b in p["blocks"] if b["type"] == "para")
    tbls = sum(1 for p in pages for b in p["blocks"] if b["type"] == "table")
    imgs = sum(1 for p in pages for b in p["blocks"] if b["type"] == "image")
    words = sum(len(b["text"].split()) for p in pages for b in p["blocks"] if b["type"] == "para")
    return {
        "filename": original_name,
        "pages": pages,
        "stats": {"pages": len(pages), "paragraphs": paras, "tables": tbls,
                  "words": words, "images": imgs},
    }
# ------------------------------------------------------------------
# Word builder
# ------------------------------------------------------------------
def add_docx_image(doc, b):
    path = STORAGE / Path(b["file"]).name
    if not path.exists():
        return
    w_in = min(max(b["width"] / 72.0, 0.2), MAX_DOCX_WIDTH_IN)
    p = doc.add_paragraph()
    p.alignment = {"left": WD_ALIGN_PARAGRAPH.LEFT,
                   "center": WD_ALIGN_PARAGRAPH.CENTER,
                   "right": WD_ALIGN_PARAGRAPH.RIGHT}.get(b.get("align"), WD_ALIGN_PARAGRAPH.LEFT)
    p.add_run().add_picture(str(path), width=Inches(w_in))
def build_docx(data):
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    for s in doc.sections:
        s.left_margin = s.right_margin = Inches(0.9)
        s.top_margin = s.bottom_margin = Inches(0.8)
    first = True
    for page in data["pages"]:
        if not page["blocks"]:
            continue
        if not first:
            doc.add_page_break()
        first = False
        for b in page["blocks"]:
            if b["type"] == "para":
                if b["heading"]:
                    doc.add_heading(b["text"], level=min(b["heading"], 3))
                else:
                    p = doc.add_paragraph()
                    run = p.add_run(b["text"])
                    run.bold = b["bold"]
            elif b["type"] == "image":
                try:
                    add_docx_image(doc, b)
                except Exception:
                    pass  
            else:
                rows = b["rows"]
                ncols = max(len(r) for r in rows)
                table = doc.add_table(rows=len(rows), cols=ncols)
                table.style = "Table Grid"
                for ri, row in enumerate(rows):
                    for ci in range(ncols):
                        cell = table.cell(ri, ci)
                        cell.text = row[ci] if ci < len(row) else ""
                        if ri == 0 and len(rows) > 1:
                            for par in cell.paragraphs:
                                for r in par.runs:
                                    r.bold = True
                doc.add_paragraph()
    if first:
        doc.add_paragraph("(No content found)")

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf
# ------------------------------------------------------------------
# Excel builder
# ------------------------------------------------------------------
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEAD_FILL = PatternFill("solid", fgColor="4F46E5")
INT_RE = re.compile(r"-?(0|[1-9]\d{0,14})")
FLOAT_RE = re.compile(r"-?(0|[1-9]\d{0,14})\.\d+")
COMMA_RE = re.compile(r"-?\d{1,3}(,\d{3})+(\.\d+)?")
def to_value(s):
    t = s.strip()
    if COMMA_RE.fullmatch(t):
        n = t.replace(",", "")
        return float(n) if "." in n else int(n)
    if INT_RE.fullmatch(t):
        return int(t)
    if FLOAT_RE.fullmatch(t):
        return float(t)
    return s
def set_cell(ws, r, c, val):
    cell = ws.cell(row=r, column=c)
    cell.value = val
    if isinstance(val, str) and val.startswith("="):
        cell.data_type = "s"  
    return cell
def write_table(ws, start_row, rows):
    ncols = max(len(r) for r in rows)
    for ri, row in enumerate(rows):
        for ci in range(ncols):
            raw = row[ci] if ci < len(row) else ""
            cell = set_cell(ws, start_row + ri, ci + 1, to_value(raw))
            cell.border = BORDER
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if ri == 0 and len(rows) > 1:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = HEAD_FILL
    return start_row + len(rows)
def table_widths(rows, base=14, cap=45):
    ncols = max(len(r) for r in rows)
    w = [base] * ncols
    for r in rows:
        for i, c in enumerate(r):
            longest = max((len(x) for x in c.split("\n")), default=0)
            w[i] = min(cap, max(w[i], longest + 2))
    return w
def add_xlsx_image(ws, r, b, max_px):
    path = STORAGE / Path(b["file"]).name
    if not path.exists():
        return r
    w_px = b["width"] * 96.0 / 72.0
    h_px = b["height"] * 96.0 / 72.0
    if w_px > max_px:  
        h_px *= max_px / w_px
        w_px = max_px
    img = XLImage(io.BytesIO(path.read_bytes()))
    img.width, img.height = max(1, int(w_px)), max(1, int(h_px))
    ws.add_image(img, f"A{r}")
    return r + math.ceil(h_px / 20.0) + 1  # default row height ~20px
def build_xlsx(data):
    wb = Workbook()
    ws = wb.active
    ws.title = "Full Document"
    all_tables = [(p["number"], b["rows"]) for p in data["pages"]
                  for b in p["blocks"] if b["type"] == "table"]
    span = max([6] + [max(len(r) for r in rows) for _, rows in all_tables])
    widths = [18] * span
    for _, rows in all_tables:
        for i, w in enumerate(table_widths(rows)):
            widths[i] = max(widths[i], w) if widths[i] != 18 else w
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    total_chars = sum(widths)
    r = 1
    for page in data["pages"]:
        if not page["blocks"]:
            continue
        c = set_cell(ws, r, 1, f"--- Page {page['number']} ---")
        c.font = Font(italic=True, color="888888")
        r += 1
        for b in page["blocks"]:
            if b["type"] == "para":
                ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=span)
                size = b["size"] or 11
                cell = set_cell(ws, r, 1, b["text"])
                cell.alignment = Alignment(wrap_text=True, vertical="top")
                cell.font = Font(bold=bool(b["heading"] or b["bold"]),
                                 size=max(10, min(size, 24)))
                lines = max(1, math.ceil(len(b["text"]) * (size / 11) / total_chars))
                ws.row_dimensions[r].height = min(409, 15 * lines * max(1, size / 11))
                r += 1
            elif b["type"] == "image":
                try:
                    r = add_xlsx_image(ws, r, b, max(60, total_chars * 7))
                except Exception:
                    pass  
            else:
                r = write_table(ws, r, b["rows"]) + 1
        r += 1
    for n, (pno, rows) in enumerate(all_tables, 1):
        sheet = wb.create_sheet(f"P{pno}_Table{n}"[:31])
        for i, w in enumerate(table_widths(rows), 1):
            sheet.column_dimensions[get_column_letter(i)].width = w
        write_table(sheet, 1, rows)
        if len(rows) > 1:
            sheet.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf
# ------------------------------------------------------------------
# Storage / cleanup / preview
# ------------------------------------------------------------------
def cleanup_old(max_age=3600):
    now = time.time()
    for f in STORAGE.glob("*"):  
        try:
            if now - f.stat().st_mtime > max_age:
                f.unlink()
        except OSError:
            pass
def build_preview(data, limit=6):
    out = []
    for p in data["pages"]:
        for b in p["blocks"]:
            if len(out) >= limit:
                return out
            if b["type"] == "para":
                out.append({"type": "para", "text": b["text"][:240],
                            "heading": b["heading"], "bold": b["bold"]})
            elif b["type"] == "table":
                out.append({"type": "table", "rows": [r[:8] for r in b["rows"][:5]]})
    return out
def build_export(fid, fmt):
    if not FID_RE.match(fid or "") or fmt not in ("word", "excel"):
        raise ValueError("Invalid request.")
    path = STORAGE / f"{fid}.json"
    if not path.exists():
        raise FileNotFoundError("Session expired. Please upload the PDF again.")
    data = json.loads(path.read_text(encoding="utf-8"))
    stem = re.sub(r"[^\w\- ]", "_", Path(data["filename"]).stem) or "converted"
    if fmt == "word":
        return build_docx(data), f"{stem}.docx", MIME_DOCX
    return build_xlsx(data), f"{stem}.xlsx", MIME_XLSX
# ------------------------------------------------------------------
# Routes
# ------------------------------------------------------------------
@app.get("/")
def index():
    return render_template("index.html")
@app.get("/favicon.ico")
def favicon():
    ico = resource_path("convert.ico")
    if ico.exists():
        return send_from_directory(str(ico.parent), ico.name, mimetype="image/x-icon")
    return ("", 204)
@app.post("/api/upload")
def upload():
    cleanup_old()
    f = request.files.get("file")
    if not f or not f.filename:
        return jsonify(error="Please select a PDF file."), 400
    if not f.filename.lower().endswith(".pdf"):
        return jsonify(error="Only PDF files are allowed."), 400
    head = f.stream.read(5)
    f.stream.seek(0)
    if head != b"%PDF-":
        return jsonify(error="This file is not a valid PDF."), 400
    fid = uuid.uuid4().hex
    pdf_path = STORAGE / f"{fid}.pdf"
    f.save(pdf_path)
    try:
        data = extract_pdf(pdf_path, f.filename, fid)
    except Exception as e:
        remove_images(fid)
        msg = str(e).lower()
        if "password" in msg or "password" in type(e).__name__.lower():
            return jsonify(error="This PDF is password protected."), 422
        return jsonify(error="Could not read this PDF. It may be corrupted."), 422
    finally:
        try:
            pdf_path.unlink()
        except OSError:
            pass
    st = data["stats"]
    if st["paragraphs"] == 0 and st["tables"] == 0:
        remove_images(fid)
        return jsonify(error="No text found. This looks like a scanned/image PDF (OCR required)."), 422
    (STORAGE / f"{fid}.json").write_text(json.dumps(data), encoding="utf-8")
    return jsonify(id=fid, filename=f.filename, stats=st, preview=build_preview(data))
@app.get("/api/download/<fid>/<fmt>")
def download(fid, fmt):
    try:
        buf, name, mime = build_export(fid, fmt)
    except ValueError as e:
        return jsonify(error=str(e)), 400
    except FileNotFoundError as e:
        return jsonify(error=str(e)), 404
    return send_file(buf, as_attachment=True, download_name=name, mimetype=mime)
@app.post("/api/discard/<fid>")
def discard(fid):
    if FID_RE.match(fid):
        try:
            (STORAGE / f"{fid}.json").unlink()
        except OSError:
            pass
        remove_images(fid)
    return jsonify(ok=True)
@app.errorhandler(413)
def too_large(_):
    return jsonify(error="File too large. Maximum size is 50 MB."), 413
if __name__ == '__main__':
    app.run(debug=True, port=5010, host='0.0.0.0')